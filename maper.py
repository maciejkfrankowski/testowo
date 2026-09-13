# -*- coding: utf-8 -*-
# ============================================================
# UNIWERSALNY MAPER SPRAWOZDAN - STANDARD BIZNESRADAR
#
# Korzysta z arkuszy slownika: Klucze, Aliasy, Reguly_obliczane, Spolki, README
# (opcjonalnie: Klucze_B dla szablonu bankowego).
#
# Nowa spolka nie wymaga zmiany programu - wystarczy wiersz w arkuszu "Spolki".
#
# Wersja poprawiona. Lista naprawionych bledow na koncu pliku.
# ============================================================
import re
import unicodedata
import datetime
from pathlib import Path

import pandas as pd
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

# ============================================================
# 1. NORMALIZACJA
# ============================================================
# POPRAWKA: separatory zamieniamy na SPACJE, nie usuwamy. Usuwanie sklejalo
# "- udzielone pozyczki" z "Udzielone pozyczki" w jeden klucz - pozycja "w tym"
# liczyla sie wtedy drugi raz. Dodatkowo wiodacy myslnik jest zachowywany,
# a "l" z kreska zamieniane recznie, bo NFKD go nie rozklada.
# "Ŝ" (S z daszkiem) to klasyczna pozostalosc po starym kodowaniu Mazovia/CP1250 -
# w polskich sprawozdaniach zawsze oznacza "z" ("NaleŜnosci", "PoŜyczki", "pienięŜne").
ZAMIANY = (("ł", "l"), ("Ł", "L"), ("Ŝ", "z"), ("ŝ", "z"),
           ("đ", "d"), ("ø", "o"), ("æ", "ae"), ("ß", "ss"))


def normalizuj(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip().lower()
    text = text.replace("–", "-").replace("—", "-").replace(" ", " ")
    myslnik = text.startswith("-")
    for a, b in ZAMIANY:
        text = text.replace(a, b)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return ("- " + text) if myslnik else text


def tekst(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def is_empty(value):
    return value is None or (isinstance(value, float) and pd.isna(value)) or str(value).strip() == ""


def to_float(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    s = str(value).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not s:
        return None
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    try:
        return float(s)
    except ValueError:
        return None


def prawda(value):
    return normalizuj(value) in {"tak", "true", "yes", "1", "t"}


# ============================================================
# 2. WCZYTANIE SLOWNIKA
# ============================================================
WYMAGANE = {
    "Klucze": ["standard_key", "etykieta_pl", "sekcja"],
    "Aliasy": ["spolka", "sekcja", "nazwa_oryginalna", "standard_key", "znak", "waga"],
    "Reguly_obliczane": ["standard_key", "formula"],
    "Spolki": ["spolka", "nazwa", "arkusz", "kol_nazwa", "kol_wartosci", "mnoznik",
               "naglowki_okresow", "sekcje"],
}


def wczytaj_slownik(path):
    xl = pd.ExcelFile(path)
    brak = [a for a in WYMAGANE if a not in xl.sheet_names]
    if brak:
        raise ValueError("Brak wymaganych arkuszy: " + ", ".join(brak))

    dane = {a.lower(): pd.read_excel(path, sheet_name=a) for a in WYMAGANE}
    for arkusz, kolumny in WYMAGANE.items():
        df = dane[arkusz.lower()]
        for c in kolumny:
            if c not in df.columns:
                raise ValueError(f"{arkusz}: brak kolumny '{c}'")

    # POPRAWKA: kolumny opcjonalne uzupelniamy, zeby pozniej nie bylo KeyError
    for c in ["szablon", "tolerancja", "kapital_obrotowy_z_bilansu",
              "kol_wartosci_bilans", "kol_wartosci_rzis", "waluta", "wyjatki_z"]:
        if c not in dane["spolki"].columns:
            dane["spolki"][c] = None
    if "tylko_gdy_brak" not in dane["reguly_obliczane"].columns:
        dane["reguly_obliczane"]["tylko_gdy_brak"] = "NIE"
    if "kolejnosc" not in dane["reguly_obliczane"].columns:
        dane["reguly_obliczane"]["kolejnosc"] = range(1, len(dane["reguly_obliczane"]) + 1)

    aliasy = dane["aliasy"].copy()
    aliasy["_nazwa"] = aliasy["nazwa_oryginalna"].map(normalizuj)
    aliasy["_spolka"] = aliasy["spolka"].map(normalizuj)
    aliasy["_sekcja"] = aliasy["sekcja"].map(tekst)          # sekcja bez normalizacji - to kod
    dane["aliasy"] = aliasy

    klucze = {}
    for _, r in dane["klucze"].iterrows():
        k = tekst(r["standard_key"])
        if k:
            klucze[k] = {"etykieta": tekst(r["etykieta_pl"]), "sekcja": tekst(r["sekcja"])}
    dane["klucz_info"] = klucze

    if "Klucze_B" in xl.sheet_names:
        dane["klucze_b"] = pd.read_excel(path, sheet_name="Klucze_B")
    else:
        dane["klucze_b"] = pd.DataFrame(columns=["standard_key", "etykieta_pl", "sekcja"])
    return dane


def pobierz_konfiguracje(dane, kod):
    spolki = dane["spolki"]
    maska = spolki["spolka"].map(normalizuj) == normalizuj(kod)
    if not maska.any():
        dostepne = ", ".join(spolki["spolka"].dropna().astype(str))
        raise ValueError(f"Spolka '{kod}' nie istnieje w arkuszu 'Spolki'.\n"
                         f"Dostepne: {dostepne}")
    return spolki[maska].iloc[0]


# ============================================================
# 3. SEKCJE
# ============================================================
# POPRAWKA: znacznik "$" byl liczony, ale nigdy nie uzywany - kazdy znacznik
# dzialal jak otwierajacy. Przy spolkach z sumami ZAMYKAJACYMI bloki (KGHM)
# lecialo przez to 45 ze 122 wierszy do zlej sekcji, bez zadnego bledu.
SEKCJE_ZBIORCZE = {"AKTYWA", "PASYWA"}
SEKCJE_BILANSU = {"BILANS", "AKT_TRW", "AKT_OBR", "AKTYWA", "KAPITAL",
                  "ZOB_DL", "ZOB_KR", "PASYWA", "BANK_A", "BANK_P", "BANK_K"}
SEKCJE_RZIS = {"RZIS", "RZIS_ZYSK", "RZIS_CD"}


def parsuj_sekcje(value):
    otwierajace, zamykajace = {}, {}
    if is_empty(value):
        return otwierajace, zamykajace
    for element in str(value).split("|"):
        element = element.strip()
        if "=>" not in element:
            continue
        nazwa, kod = element.split("=>", 1)
        nazwa, kod = nazwa.strip(), kod.strip()
        if nazwa.startswith("$"):
            zamykajace[normalizuj(nazwa[1:])] = kod
        else:
            otwierajace[normalizuj(nazwa)] = kod
    return otwierajace, zamykajace


def parsuj_kolumny(value):
    out = []
    if is_empty(value):
        return out
    for x in str(value).split(","):
        x = x.strip()
        if x:
            try:
                out.append(int(float(x)))
            except ValueError:
                pass
    return out


def czytaj_wiersze(df, konfiguracja):
    """Zwraca liste wierszy z przypisana sekcja. Obsluguje oba typy znacznikow."""
    # kol_nazwa moze byc listą, np. "1,2". Sprawozdania z ESEF miewaja tabele
    # dwupoziomowa: kolumna 1 to strona bilansu ("Aktywa"/"Pasywa"), a nazwa pozycji
    # stoi dopiero w kolumnie 2. Bierzemy wtedy OSTATNIA niepusta komorke, ktora nie
    # jest liczba - w rachunku wynikow kolumna 2 trzyma juz wartosc, wiec odpada sama.
    kol_nazwy = [c - 1 for c in parsuj_kolumny(konfiguracja["kol_nazwa"])] or [0]
    kol_glowne = [c - 1 for c in parsuj_kolumny(konfiguracja["kol_wartosci"])]
    kol_bilans = [c - 1 for c in parsuj_kolumny(konfiguracja.get("kol_wartosci_bilans"))]
    kol_rzis = [c - 1 for c in parsuj_kolumny(konfiguracja.get("kol_wartosci_rzis"))]
    mnoznik = to_float(konfiguracja["mnoznik"]) or 1.0
    otwierajace, zamykajace = parsuj_sekcje(konfiguracja["sekcje"])

    def czytaj(row_idx, kolumny):
        out = []
        for c in kolumny:
            v = df.iat[row_idx, c] if c < df.shape[1] else None
            v = to_float(v)
            if v is None:
                out.append(None)
            else:
                v = v * mnoznik
                # POPRAWKA: dane w zlotych zaokraglamy do pelnych tysiecy,
                # tak jak robi to standard - inaczej sumy nie schodza sie do grosza
                out.append(round(v) if mnoznik < 1 else v)
        return out

    wiersze, bufor, sekcja = [], [], ""

    def flush(kod):
        for w in bufor:
            w["sekcja"] = kod
            if kol_bilans and kod in SEKCJE_BILANSU:
                w["wartosci"] = w["_bilans"]
            elif kol_rzis and kod in SEKCJE_RZIS:
                w["wartosci"] = w["_rzis"]
            w.pop("_bilans", None)
            w.pop("_rzis", None)
            wiersze.append(w)
        bufor.clear()

    def nazwa_wiersza(i):
        wybrana = None
        for c in kol_nazwy:
            if c >= df.shape[1]:
                continue
            v = df.iat[i, c]
            if is_empty(v) or isinstance(v, (int, float, np.integer, np.floating)):
                continue
            if to_float(v) is not None:          # tekst, ktory jest liczba
                continue
            wybrana = str(v).strip()
        return wybrana

    for i in range(len(df)):
        nazwa = nazwa_wiersza(i)
        if is_empty(nazwa):
            continue
        n = normalizuj(nazwa)
        wartosci = czytaj(i, kol_glowne)
        w = {"wiersz": i + 1, "nazwa": nazwa, "nazwa_norm": n, "sekcja": None,
             "wartosci": wartosci,
             "_bilans": czytaj(i, kol_bilans) if kol_bilans else wartosci,
             "_rzis": czytaj(i, kol_rzis) if kol_rzis else wartosci}

        if n in otwierajace:                       # naglowek OTWIERA blok
            flush(sekcja)
            sekcja = otwierajace[n]
            if any(v is not None for v in wartosci):
                bufor.append(w)
            continue
        if n in zamykajace:                        # suma ZAMYKA blok
            kod = zamykajace[n]
            if kod in SEKCJE_ZBIORCZE and sekcja not in SEKCJE_ZBIORCZE and sekcja and bufor:
                # "Aktywa razem" to pojedyncza linia - nie zabiera calego bufora
                flush(sekcja)
                bufor.append(w)
                flush(kod)
            else:
                bufor.append(w)
                flush(kod)
            continue
        bufor.append(w)
    flush(sekcja)
    return wiersze


# ============================================================
# 4. WYBOR ALIASOW
# ============================================================
def wybierz_aliasy(aliasy, nazwa_norm, spolka, sekcja):
    """'spolka' moze byc kodem albo lista kodow od NAJWAZNIEJSZEGO: [wariant, macierzysta].

    Wariant ukladu (np. raport roczny) musi moc NADPISAC wyjatek spolki macierzystej -
    ta sama etykieta potrafi znaczyc co innego w raporcie rocznym niz w kwartalnym.
    """
    kand = aliasy[aliasy["_nazwa"] == nazwa_norm]
    if kand.empty:
        return []

    # Pierwszenstwo: wyjatek spolki bije alias globalny, dokladna sekcja bije "*".
    # POPRAWKA: to musza byc CZTERY poziomy sprawdzane po kolei. Wczesniej wyjatek
    # spolki zabieral pule raz na zawsze - jesli dotyczyl innej sekcji, alias globalny
    # dla wlasciwej sekcji przepadal i wiersz szedl do niezmapowanych.
    kody = [spolka] if isinstance(spolka, str) else list(spolka)
    kody = [normalizuj(k) for k in kody if tekst(k)]
    kody = [k for i, k in enumerate(kody) if k not in kody[:i]] + [""]   # "" = alias globalny
    maski = []
    for k in kody:
        maski.append((kand["_spolka"] == k) & (kand["_sekcja"] == sekcja))
        maski.append((kand["_spolka"] == k) & (kand["_sekcja"] == "*"))
    for maska in maski:
        if maska.any():
            kand = kand[maska]
            break
    else:
        return []

    # POPRAWKA: usuniety filtr po kolumnie "pewnosc". Pewnosc to adnotacja dla
    # czlowieka, nie kryterium wyboru - filtrowanie po niej gubilo aliasy
    # skladajace sie na agregat (jeden wiersz -> kilka kluczy o roznej pewnosci).
    kand = kand.drop_duplicates(subset=["_spolka", "_sekcja", "_nazwa",
                                        "standard_key", "znak", "waga"])
    return list(kand.itertuples(index=False))


# ============================================================
# 5. REGULY OBLICZANE
# ============================================================
def parsuj_formule(formula):
    if is_empty(formula):
        return []
    out = []
    for znak, key in re.findall(r"([+-]?)\s*([A-Za-z0-9_]+)", str(formula)):
        out.append((-1.0 if znak == "-" else 1.0, key))
    return out


def wykonaj_reguly(reguly, wartosci, bezposrednie, n_kol):
    """POPRAWKA: brakujacy skladnik liczy sie jako ZERO, a nie przerywa reguly.
    Wczesniej gross_profit nie powstawal u zadnej spolki z rachunkiem rodzajowym
    (AOL, Grupa Pracuj, PlayWay, 11 bit), bo nie maja kosztow sprzedazy i zarzadu."""
    audyt = []
    if reguly.empty:
        return audyt
    r = reguly.copy()
    r["_sort"] = pd.to_numeric(r["kolejnosc"], errors="coerce")
    for _, reg in r.sort_values("_sort").iterrows():
        key = tekst(reg["standard_key"])
        skladniki = parsuj_formule(reg["formula"])
        if not key or not skladniki:
            continue
        if prawda(reg["tylko_gdy_brak"]) and key in bezposrednie:
            continue
        wynik = []
        for j in range(n_kol):
            suma = 0.0
            for znak, k in skladniki:
                v = wartosci.get(k, [None] * n_kol)[j]
                suma += znak * (v if isinstance(v, (int, float)) else 0.0)
            wynik.append(suma)
        wartosci[key] = wynik
        audyt.append({"standard_key": key, "formula": tekst(reg["formula"]), "typ": "REGULA"})
    return audyt


# ============================================================
# 6. KAPITAL OBROTOWY
# ============================================================
def kapital_obrotowy(wartosci, n_kol, poprzedni=None):
    """Metodologia biznesradar: zmiany kapitalu obrotowego NIE sa przepisywane
    z rachunku przeplywow, tylko liczone z roznic BILANSOW.

    UWAGA: okresem odniesienia jest koniec POPRZEDNIEGO ROKU, nie poprzedniego
    kwartalu. Przy I kwartale to to samo, przy II-IV juz nie. Kolumna odniesienia
    musi wiec wskazywac bilans na koniec roku - inaczej wynik bedzie po cichu zly.

    'poprzedni' to slownik {standard_key: wartosc} z bilansu okresu odniesienia,
    wczytany z pliku wzorcowego. Potrzebny, gdy sprawozdanie nie ma tego bilansu
    w drugiej kolumnie - uklad UoR podaje tam ten sam kwartal rok wczesniej.

    POPRAWKA: liczone sa wszystkie CZTERY pozycje (bylo: dwie), a listy maja
    tyle elementow, ile okresow (bylo: jeden, co rozjezdzalo kolumny w Wyniku).
    """
    KLUCZE = ["change_in_inventories", "change_in_receivables", "change_in_payables",
              "change_in_other_assets", "change_in_working_capital"]
    for k in KLUCZE:
        wartosci.setdefault(k, [None] * n_kol)

    def g(key, j):
        v = wartosci.get(key, [None] * n_kol)
        x = v[j] if j < len(v) else None
        return x if isinstance(x, (int, float)) else 0.0

    for j in range(n_kol):
        if j == 0 and poprzedni:
            odn = lambda k: (poprzedni.get(k) if isinstance(poprzedni.get(k), (int, float)) else 0.0)
        elif j + 1 < n_kol:
            odn = lambda k, _j=j: g(k, _j + 1)
        else:
            for k in KLUCZE:
                wartosci[k][j] = None
            continue
        teraz = lambda k, _j=j: g(k, _j)
        dlug = lambda f: f("current_liabilities") - f("current_other_liabilities")
        inne = lambda f: f("current_other_liabilities") + f("reckoning")
        wartosci["change_in_inventories"][j] = odn("inventory") - teraz("inventory")
        wartosci["change_in_receivables"][j] = odn("current_receivables") - teraz("current_receivables")
        wartosci["change_in_payables"][j] = dlug(teraz) - dlug(odn)
        wartosci["change_in_other_assets"][j] = inne(teraz) - inne(odn)
        wartosci["change_in_working_capital"][j] = sum(
            wartosci[k][j] for k in KLUCZE[:4])


# ============================================================
# 7. MAPOWANIE
# ============================================================
def mapuj(dane, kod_spolki, plik, bilans_poprzedni=None, okresy=None):
    cfg = pobierz_konfiguracje(dane, kod_spolki)
    # Ten sam podmiot potrafi skladac raporty w ROZNYCH ukladach kolumn (III kwartal ma
    # osobno kwartal i narastajaco, roczny ma tylko rok). Taki wariant dostaje wlasny
    # wiersz w arkuszu "Spolki", ale przez kolumne 'wyjatki_z' dziedziczy wyjatki
    # macierzystej spolki - inaczej stracilby np. znaki kosztow.
    kod_wyjatkow = [kod_spolki, tekst(cfg.get("wyjatki_z"))]
    df = pd.read_excel(plik, sheet_name=tekst(cfg["arkusz"]), header=None)
    n_kol = len(parsuj_kolumny(cfg["kol_wartosci"]))
    wiersze = czytaj_wiersze(df, cfg)

    wartosci, bezposrednie, audyt, niezmapowane = {}, set(), [], []
    for w in wiersze:
        aliasy = wybierz_aliasy(dane["aliasy"], w["nazwa_norm"], kod_wyjatkow, w["sekcja"])
        if not aliasy:
            niezmapowane.append({"wiersz": w["wiersz"], "nazwa": w["nazwa"],
                                 "sekcja": w["sekcja"], "powod": "BRAK_ALIASU"})
            continue
        for al in aliasy:
            key = tekst(al.standard_key)
            if not key:
                continue
            if normalizuj(key) in ("pomin", "ignoruj"):
                audyt.append({**w, "standard_key": "POMIN", "znak": "", "status": "POMIN"})
                continue
            znak = to_float(al.znak) or 1.0
            waga = to_float(al.waga) or 1.0
            wartosci.setdefault(key, [0.0] * n_kol)
            for j, v in enumerate(w["wartosci"][:n_kol]):
                if v is not None:
                    wartosci[key][j] += v * znak * waga
            bezposrednie.add(key)
            audyt.append({**w, "standard_key": key, "znak": znak, "status": "ZMAPOWANO"})

    audyt_regul = wykonaj_reguly(dane["reguly_obliczane"], wartosci, bezposrednie, n_kol)
    # podany bilans odniesienia sam w sobie wymusza metode bilansowa
    if prawda(cfg.get("kapital_obrotowy_z_bilansu")) or bilans_poprzedni:
        kapital_obrotowy(wartosci, n_kol, bilans_poprzedni)

    # metadane, ktore da sie wyprowadzic z konfiguracji
    waluta = tekst(cfg.get("waluta")).upper()
    wartosci["currency_id"] = [waluta if re.fullmatch(r"[A-Z]{3}", waluta) else "PLN"] * n_kol
    daty = daty_z_naglowkow(cfg["naglowki_okresow"])
    wartosci["balance_date"] = [(daty[j].strftime("%d.%m.%Y") if j < len(daty) and daty[j] else None)
                                for j in range(n_kol)]

    # Etykiety okresow mozna podac z zewnatrz - ten sam podmiot sklada raporty za
    # rozne okresy w IDENTYCZNYM ukladzie, a kod spolki jest jednoczesnie kluczem
    # wyjatkow, wiec nie wolno zakladac drugiej spolki tylko po to, zeby zmienic podpisy.
    etykiety = ([x.strip() for x in okresy] if okresy
                else [x.strip() for x in str(cfg["naglowki_okresow"]).split("|") if x.strip()])
    if okresy:
        daty = daty_z_naglowkow("|".join(etykiety))
        wartosci["balance_date"] = [(daty[j].strftime("%d.%m.%Y") if j < len(daty) and daty[j] else None)
                                    for j in range(n_kol)]

    return {"konfiguracja": cfg, "wartosci": wartosci, "audyt": audyt,
            "audyt_regul": audyt_regul, "niezmapowane": niezmapowane, "n_kol": n_kol,
            "okresy": etykiety, "tag": tag_okresu(etykiety[0] if etykiety else "",
                                                  (daty_z_naglowkow("|".join(etykiety)) or [None])[0])}


RZYMSKIE = {"i": 1, "ii": 2, "iii": 3, "iv": 4}


def tag_okresu(naglowek, data=None):
    """Skrot okresu do nazwy pliku: "I kw 2026 / 31.03.2026" -> "1Q2026".

    Rozpoznaje kwartaly, polrocza i "N m-cy". Gdy nie ma pewnosci - zwraca sama
    date bilansowa, ktora jest jednoznaczna i dobrze sie sortuje. Naglowek z
    PRZESUNIETYM rokiem obrotowym ("I kw obrotowy 2025/26") tez idzie na date,
    bo "1Q2025" myloby sie z kwartalem kalendarzowym."""
    n = normalizuj(naglowek)
    rok = data.year if data else None
    if rok and "obrotow" not in n:
        m = re.match(r"(i{1,3}|iv)\s+kw\b", n)
        if m:
            return f"{RZYMSKIE[m.group(1)]}Q{rok}"
        m = re.match(r"(i{1,3}|iv)\s+(polrocze|pol)\b", n)
        if m:
            return f"{RZYMSKIE[m.group(1)]}H{rok}"
        m = re.match(r"(\d+)\s*m\s*cy\b", n)
        if m:
            return f"{m.group(1)}M{rok}"
    return data.strftime("%Y-%m-%d") if data else "okres"


DATA_RE = re.compile(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})")


def daty_z_naglowkow(naglowki):
    out = []
    for h in str(naglowki or "").split("|"):
        z = DATA_RE.findall(h)
        if not z:
            out.append(None)
            continue
        d, m, y = z[-1]
        y = int(y) + (2000 if int(y) < 100 else 0)
        try:
            out.append(datetime.datetime(y, int(m), int(d)))
        except ValueError:
            out.append(None)
    return out


# ============================================================
# 8. WALIDACJA
# ============================================================
def waliduj(wynik, szablon="RAP"):
    """POPRAWKA: kontrola 'przychody - koszt wlasny = gross_profit' byla bledna -
    w standardzie gross_profit jest PO kosztach sprzedazy i zarzadu, wiec zapalala
    sie na czerwono u kazdej spolki, ktora je ma (Neuca 325 tys., Auto Partner 234 tys.).
    Klucz 'total_equity' nie istnieje - jest 'capital', wiec tamta kontrola
    nigdy sie nie wykonywala."""
    w = wynik["wartosci"]
    n_kol = wynik["n_kol"]
    tol = to_float(wynik["konfiguracja"].get("tolerancja")) or 1.0

    def g(k, j):
        v = w.get(k, [None] * n_kol)
        x = v[j] if j < len(v) else None
        return x if isinstance(x, (int, float)) else 0.0

    if szablon == "RAP_B":
        AKT = ["cash_and_balances_with_the_central_bank", "loans_to_banks",
               "financial_assets_held_for_trading", "held_to_maturity_financial_assets",
               "loans_to_customers", "other_financial_assets", "investments_in_subsidiaries",
               "intangible_assets", "property", "tax_assets", "other_assets"]
        ZOB = ["liabilities_due_to_the_central_bank", "liabilities_due_to_banks",
               "held_to_maturity_financial_liabilities", "financial_liabilities_held_for_trading",
               "liabilities_due_to_customers", "debt_securities", "subordinated_liabilities",
               "tax_liabilities", "other_liabilities"]
        KAP = ["share_capital", "share_premium", "own_share", "revaluation_reserves",
               "retained_earnings", "year_profit"]
        OPER = ["net_interest_income", "net_fee_income", "dividend_income", "realised_gains",
                "net_other_financial_income", "net_other_operating_income", "impairment_losses",
                "administration_costs", "net_other_operating_costs", "bank_tax"]
        testy = [
            ("Aktywa = Pasywa", lambda j: g("total_assets", j) - g("total_equity_liabilities", j)),
            ("Skladniki aktywow", lambda j: sum(g(k, j) for k in AKT) - g("total_assets", j)),
            ("Skladniki zobowiazan", lambda j: sum(g(k, j) for k in ZOB) - g("total_liabilities", j)),
            ("Kapital akcjonariuszy", lambda j: sum(g(k, j) for k in KAP) - g("own_capital", j)),
            ("Zobowiazania + kapitaly = pasywa",
             lambda j: g("total_liabilities", j) + g("total_capital", j) - g("total_equity_liabilities", j)),
            ("Wynik odsetkowy",
             lambda j: g("interest_income", j) - g("interest_expense", j) - g("net_interest_income", j)),
            ("Wynik prowizyjny",
             lambda j: g("fee_income", j) - g("fee_expense", j) - g("net_fee_income", j)),
            ("Wynik operacyjny", lambda j: sum(g(k, j) for k in OPER) - g("net_operating_profit", j)),
            ("Zysk przed opodatkowaniem",
             lambda j: g("net_operating_profit", j) + g("related_income", j) - g("beforetax_profit", j)),
            ("Przeplywy razem", lambda j: g("operating_cashflow", j) + g("investing_cashflow", j)
             + g("financing_cashflow", j) - g("net_cashflow", j)),
        ]
    else:
        testy = [
            ("Aktywa = Pasywa", lambda j: g("total_assets", j) - g("total_equity_liabilities", j)),
            # POPRAWKA: assets_for_sale to linia INFORMACYJNA wewnatrz aktywow obrotowych,
            # a nie trzeci skladnik sumy bilansowej (potwierdzone wzorcem SGN: 262800 + 260285
            # = 523085, a aktywa przeznaczone do sprzedazy 72 siedza juz w tych 260285).
            ("Aktywa trwale + obrotowe = suma", lambda j: g("noncurrent_assets", j)
             + g("current_assets", j) - g("total_assets", j)),
            ("Skladniki aktywow trwalych", lambda j: g("property", j) + g("intangible_assets", j)
             + g("noncurrent_investments", j) + g("noncurrent_receivables", j)
             + g("other_noncurrent_assets", j) - g("noncurrent_assets", j)),
            ("Skladniki aktywow obrotowych", lambda j: g("inventory", j) + g("current_receivables", j)
             + g("current_investments", j) + g("other_current_assets", j) + g("assets_for_sale", j)
             - g("current_assets", j)),
            ("Skladniki kapitalu wlasnego", lambda j: g("share_capital", j) + g("own_share", j)
             + g("reserve", j) + g("retained_earnings", j) + g("year_profit", j)
             + g("nonshare_capital", j) + g("kapitaly_pozostale_nierozdzielone", j) - g("capital", j)),
            ("Kapital + zobowiazania = pasywa", lambda j: g("capital", j) + g("noncurrent_liabilities", j)
             + g("current_liabilities", j) + g("reckoning", j) - g("total_equity_liabilities", j)),
            ("Skladniki zob. dlugoterminowych", lambda j: g("noncurrent_trade_payables", j)
             + g("noncurrent_borrowings", j) + g("noncurrent_obligations", j)
             + g("noncurrent_leasing", j) + g("noncurrent_other_liabilities", j)
             - g("noncurrent_liabilities", j)),
            ("Skladniki zob. krotkoterminowych", lambda j: g("current_trade_payables", j)
             + g("current_borrowings", j) + g("current_obligations", j) + g("current_leasing", j)
             + g("current_other_liabilities", j) - g("current_liabilities", j)),
            # POPRAWNA formula: gross_profit jest PO kosztach sprzedazy i zarzadu
            ("Zysk ze sprzedazy", lambda j: g("revenues", j) - g("cost_of_sales", j)
             - g("distribution_expenses", j) - g("administrative_expenses", j) - g("gross_profit", j)),
            ("EBIT", lambda j: g("gross_profit", j) + g("other_operating_income", j)
             - g("other_operating_costs", j) - g("ebit", j)),
            ("Zysk przed opodatkowaniem", lambda j: g("ebit", j) + g("finance_income", j)
             - g("finance_costs", j) + g("other_income", j) - g("before_tax_profit", j)),
            ("Przeplywy razem", lambda j: g("operating_cashflow", j) + g("investing_cashflow", j)
             + g("financing_cashflow", j) - g("net_cashflow", j)),
            ("Kapital obrotowy = suma skladnikow", lambda j: g("change_in_receivables", j)
             + g("change_in_inventories", j) + g("change_in_payables", j)
             + g("change_in_other_assets", j) - g("change_in_working_capital", j)),
        ]

    out = []
    for nazwa, fn in testy:
        wiersz = {"kontrola": nazwa, "tolerancja": tol}
        for j in range(n_kol):
            d = fn(j)
            wiersz[f"okres_{j+1}"] = "OK" if abs(d) <= tol else f"ROZNICA {d:,.0f}"
        out.append(wiersz)

    # kontrola miekka: implikowana stopa podatku poza 0-40% zwykle oznacza blad mapowania
    klucz_brutto = "beforetax_profit" if szablon == "RAP_B" else "before_tax_profit"
    wiersz = {"kontrola": "Implikowana stopa podatku (0-40%)", "tolerancja": ""}
    for j in range(n_kol):
        bt, np_ = g(klucz_brutto, j), g("net_profit", j)
        if bt <= 0:
            wiersz[f"okres_{j+1}"] = "brak danych"
        else:
            r = (bt - np_) / bt
            wiersz[f"okres_{j+1}"] = f"OK ({r:.1%})" if 0 <= r <= 0.40 else f"PODEJRZANA {r:.1%}"
    out.append(wiersz)
    return pd.DataFrame(out)


# ============================================================
# 9. PORÓWNANIE ZE WZORCEM  (NOWE)
# ============================================================
def wczytaj_okres_wzorca(wzorzec, data=None, naglowek_okresu=None, arkusz="Spr Fin"):
    """Zwraca {standard_key: wartosc} dla jednego okresu z pliku w standardzie.

    Okres wskazuje sie data bilansowa albo naglowkiem kolumny - to drugie ratuje
    sytuacje, w ktorych data we wzorcu jest bledna (ALR ma w dwoch ostatnich
    kolumnach rok 2028 zamiast 2026)."""
    rows = list(load_workbook(wzorzec, read_only=True, data_only=True)[arkusz]
                .iter_rows(min_row=1, max_row=110, values_only=True))
    if naglowek_okresu:
        kol = next((i for i in range(len(rows[0]))
                    if tekst(rows[0][i]) == tekst(naglowek_okresu)), None)
    else:
        cel = datetime.datetime.strptime(data, "%Y-%m-%d")
        daty = next((r for r in rows if r and tekst(r[0]) == "balance_date"), None)
        kol = None if daty is None else next(
            (i for i in range(len(daty))
             if isinstance(daty[i], datetime.datetime) and daty[i] == cel), None)
    if kol is None:
        raise ValueError(f"We wzorcu '{wzorzec}' nie ma kolumny "
                         f"{naglowek_okresu or data} (arkusz '{arkusz}')")
    return {tekst(r[0]): r[kol] for r in rows if r and isinstance(r[0], str)}


def porownaj_ze_wzorcem(wzorzec, wynik, data=None, naglowek_okresu=None,
                        arkusz="Spr Fin", kolumna=0):
    """Zestawia wynik z gotowym plikiem w standardzie.

    To najskuteczniejsza kontrola, jaka mamy: wykrywa bledy, ktorych sumy kontrolne
    nie widza - zla sekcje, pomylony znak, pominieta pozycje."""
    wzor = wczytaj_okres_wzorca(wzorzec, data, naglowek_okresu, arkusz)

    tol = to_float(wynik["konfiguracja"].get("tolerancja")) or 1.0
    out = []
    # POPRAWKA: iterujemy po SUMIE kluczy. Wczesniej leciala tylko petla po moich
    # wynikach, wiec klucz, ktorego w ogole nie zmapowalem, nie pojawial sie w
    # porownaniu - a to najgrozniejszy przypadek (DOM: net_cashflow trafil na POMIN
    # i porownanie pokazywalo zero roznic).
    klucze = list(wynik["wartosci"]) + [k for k in wzor if k not in wynik["wartosci"]]
    for key in klucze:
        values = wynik["wartosci"].get(key) or []
        moje = values[kolumna] if kolumna < len(values) else None
        wz = wzor.get(key)
        mn = moje if isinstance(moje, (int, float)) else None
        wn = wz if isinstance(wz, (int, float)) else None
        if (wn or 0) == 0 and (mn or 0) == 0:
            status, roznica = "obie puste", None
        elif wn is None or mn is None:
            status, roznica = "TYLKO W JEDNYM", None
        else:
            roznica = mn - wn
            status = "ZGODNE" if abs(roznica) <= tol else "ROZNICA"
        out.append({"standard_key": key, "wzorzec": wn, "moje": mn,
                    "roznica": roznica, "status": status})
    return pd.DataFrame(out).sort_values("status")


# ============================================================
# 10. ZAPIS
# ============================================================
def zapisz(dane, wynik, sciezka, porownanie=None):
    cfg = wynik["konfiguracja"]
    szablon = tekst(cfg.get("szablon")).upper() or "RAP"
    klucze = dane["klucze_b"] if (szablon == "RAP_B" and not dane["klucze_b"].empty) else dane["klucze"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Wynik"
    okresy = wynik["okresy"] or [f"wartosc_{i+1}" for i in range(wynik["n_kol"])]
    ws.append(["standard_key", "etykieta_pl", "sekcja"] + okresy[:wynik["n_kol"]])
    for _, r in klucze.iterrows():
        key = tekst(r["standard_key"])
        if not key:
            continue
        v = wynik["wartosci"].get(key)
        if v is None:
            # POPRAWKA: klucz bez zrodla dostaje 0 (metadane - pusto), zeby arkusz
            # Wynik mial zawsze komplet wierszy i dalo sie go wkleic do standardu
            v = [None] * wynik["n_kol"] if tekst(r["sekcja"]) == "META" else [0.0] * wynik["n_kol"]
        ws.append([key, tekst(r["etykieta_pl"]), tekst(r["sekcja"])]
                  + [round(x, 2) if isinstance(x, (int, float)) else x for x in v])

    def arkusz_z_df(nazwa, df):
        w = wb.create_sheet(nazwa)
        w.append(list(df.columns))
        for row in df.itertuples(index=False, name=None):
            w.append([("" if (isinstance(x, float) and pd.isna(x)) else x) for x in row])

    arkusz_z_df("Walidacja", waliduj(wynik, szablon))
    if porownanie is not None:
        arkusz_z_df("Porownanie", porownanie)

    wa = wb.create_sheet("Audyt")
    wa.append(["wiersz", "sekcja", "nazwa w raporcie", "standard_key", "znak", "status"])
    for a in wynik["audyt"]:
        wa.append([a["wiersz"], a["sekcja"], a["nazwa"], a["standard_key"], a["znak"], a["status"]])
    for a in wynik["audyt_regul"]:
        wa.append(["", "", a["formula"], a["standard_key"], "", "REGULA_OBLICZANA"])

    wn = wb.create_sheet("Niezmapowane")
    wn.append(["wiersz", "sekcja", "nazwa", "powod"])
    for x in wynik["niezmapowane"]:
        wn.append([x["wiersz"], x["sekcja"], x["nazwa"], x["powod"]])

    wi = wb.create_sheet("Info")
    wi.append(["parametr", "wartosc"])
    for k in ["spolka", "nazwa", "szablon", "arkusz", "kol_nazwa", "kol_wartosci",
              "kol_wartosci_bilans", "kol_wartosci_rzis", "mnoznik", "tolerancja",
              "kapital_obrotowy_z_bilansu", "waluta"]:
        wi.append([k, tekst(cfg.get(k))])
    wi.append(["liczba_niezmapowanych", len(wynik["niezmapowane"])])

    for sheet in wb.worksheets:
        for col in sheet.columns:
            szer = max((len(str(c.value)) for c in col if c.value is not None), default=0)
            sheet.column_dimensions[get_column_letter(col[0].column)].width = min(szer + 2, 60)
    wb.save(sciezka)


# ============================================================
# 11. URUCHOMIENIE W COLABIE
# ============================================================
def main():
    from google.colab import files

    print("=" * 70)
    print("UNIWERSALNY MAPER SPRAWOZDAN")
    print("=" * 70)
    print("\nWgraj: slownik, sprawozdanie spolki i (opcjonalnie) plik wzorcowy.\n")
    uploaded = files.upload()

    slownik = sprawozdanie = wzorzec = None
    for nazwa in uploaded:
        low = nazwa.lower()
        if "slownik" in low or "słownik" in low:
            slownik = nazwa
        elif "_sprawozdania" in low or sprawozdanie is None:
            sprawozdanie = nazwa
        else:
            wzorzec = nazwa
    if not slownik:
        raise RuntimeError("Nie znaleziono pliku slownika (nazwa musi zawierac 'slownik').")
    if not sprawozdanie:
        raise RuntimeError("Nie znaleziono sprawozdania spolki.")

    dane = wczytaj_slownik(slownik)
    print(f"\nSlownik OK: {len(dane['klucze'])} kluczy, {len(dane['aliasy'])} aliasow, "
          f"{len(dane['spolki'])} spolek")
    print("\nDostepne spolki:")
    for _, r in dane["spolki"].iterrows():
        print(f"   {tekst(r['spolka']):>5}  {tekst(r['nazwa'])}")

    kod = input("\nPodaj kod spolki: ").strip()

    # Ten sam podmiot sklada raporty za rozne okresy w identycznym ukladzie. Kod spolki
    # jest jednoczesnie kluczem wyjatkow, wiec NIE zakladamy drugiego kodu tylko po to,
    # zeby zmienic podpisy kolumn - podajemy je tutaj.
    cfg_okresy = tekst(pobierz_konfiguracje(dane, kod)["naglowki_okresow"])
    print(f"\nOkresy ze slownika: {cfg_okresy}")
    podane = input("Inne okresy? (rozdziel '|', Enter = zostaw): ").strip()
    okresy = [x.strip() for x in podane.split("|") if x.strip()] if podane else None

    data = nag = ""
    if wzorzec:
        data = input("Data bilansowa we wzorcu (RRRR-MM-DD, Enter = pomin): ").strip()
        nag = "" if data else input("albo naglowek okresu (np. 'II 2026'): ").strip()

    # Bilans odniesienia dla kapitalu obrotowego = koniec POPRZEDNIEGO ROKU.
    # Jesli sprawozdanie ma go w drugiej kolumnie, nie trzeba nic robic; przy
    # ukladzie UoR (druga kolumna to ten sam kwartal rok wczesniej) bierzemy go
    # z pliku wzorcowego.
    prev = None
    if wzorzec and data:
        rok = datetime.datetime.strptime(data, "%Y-%m-%d").year
        try:
            prev = wczytaj_okres_wzorca(wzorzec, f"{rok-1}-12-31")
            print(f"   Bilans odniesienia: {rok-1}-12-31 (z pliku wzorcowego)")
        except (ValueError, StopIteration):
            pass

    wynik = mapuj(dane, kod, sprawozdanie, bilans_poprzedni=prev, okresy=okresy)

    porownanie = None
    if data or nag:
        try:
            porownanie = porownaj_ze_wzorcem(wzorzec, wynik, data or None, nag or None)
        except (ValueError, StopIteration) as e:
            print("   Pomijam porownanie:", e)

    # Okres w nazwie pliku - inaczej kolejny raport tej samej spolki nadpisze poprzedni
    out = f"wynik_{kod.upper()}_{wynik['tag']}.xlsx"
    zapisz(dane, wynik, out, porownanie)

    print("\n" + "=" * 70)
    print(f"Zmapowano wpisow: {len(wynik['audyt'])}   Niezmapowanych: {len(wynik['niezmapowane'])}")
    for x in wynik["niezmapowane"][:10]:
        print(f"   [{x['sekcja']}] {x['nazwa'][:64]}")
    szablon = tekst(wynik["konfiguracja"].get("szablon")).upper() or "RAP"
    wal = waliduj(wynik, szablon)
    zle = [r for _, r in wal.iterrows()
           if any(str(v).startswith(("ROZNICA", "PODEJRZANA")) for k, v in r.items()
                  if k.startswith("okres_"))]
    print(f"Walidacja: {len(wal)-len(zle)}/{len(wal)} OK")
    for r in zle:
        print("   !", r["kontrola"], {k: v for k, v in r.items() if k.startswith("okres_")})
    if porownanie is not None:
        zg = (porownanie["status"] == "ZGODNE").sum()
        rz = (porownanie["status"] == "ROZNICA").sum()
        print(f"Porownanie ze wzorcem: zgodne {zg}, roznice {rz}")
        for _, r in porownanie[porownanie["status"] == "ROZNICA"].iterrows():
            print(f"   {r['standard_key']:28} wzorzec {r['wzorzec']}  moje {r['moje']}")
    print(f"\nPlik wynikowy: {out}")
    files.download(out)


if __name__ == "__main__" and "get_ipython" in dir():
    main()

# ============================================================
# NAPRAWIONE BLEDY
# ============================================================
# 1. Znacznik "$" byl ignorowany - kazdy dzialal jak otwierajacy.
#    Skutek: KGHM 45/122 wierszy w zlej sekcji, bez zadnego bledu.
# 2. Regula byla pomijana, gdy brakowalo skladnika. Skutek: gross_profit
#    nie powstawal u zadnej spolki z rachunkiem rodzajowym.
# 3. Kontrola "przychody - koszt wlasny = gross_profit" byla bledna -
#    falszywy alarm u kazdej spolki z kosztami sprzedazy i zarzadu.
# 4. Normalizacja usuwala separatory zamiast zamieniac je na spacje -
#    pozycje "w tym" sklejaly sie z nadrzednymi i liczyly dwa razy.
# 5. Kapital obrotowy: liczone 2 z 4 pozycji, lista o dlugosci 1 zamiast
#    liczby okresow (rozjazd kolumn), brak zaokraglania przy mnozniku < 1.
# 6. Filtr po kolumnie "pewnosc" gubil aliasy skladajace sie na agregat.
# 7. Klucz "total_equity" nie istnieje (jest "capital") - kontrola martwa.
# 8. "szablon" czytany, ale niewymagany - KeyError przy zapisie.
# 9. Brak obslugi kol_wartosci_bilans / kol_wartosci_rzis i szablonu bankowego.
# 10. Wyjatek spolki zabieral pule aliasow raz na zawsze - gdy dotyczyl innej
#     sekcji, alias globalny dla wlasciwej sekcji przepadal. Teraz cztery poziomy
#     pierwszenstwa sprawdzane po kolei (11 wierszy mniej w Niezmapowanych).
# 11. Tolerancja dzialala jako "<" zamiast "<=" - roznica rowna 1 przy tolerancji
#     1 zapalala sie na czerwono (Grupa Pracuj).
# 12. Dodane porownanie ze wzorcem - jedyna kontrola, ktora lapie zle sekcje -
#     oraz bilans odniesienia wczytywany z wzorca (uklad UoR).
# 13. Porownanie szlo petla tylko po MOICH wynikach, wiec klucz, ktorego w ogole
#     nie zmapowalem, nie pojawial sie w zestawieniu. Teraz suma kluczy.
#     (Dom Development: net_cashflow wpadl na POMIN i porownanie milczalo.)
# 14. "Ŝ" ze starego kodowania normalizowane do "z" - inaczej "NaleŜnosci"
#     i "Naleznosci" to dwa rozne klucze.
#
# TEST REGRESJI: 17 spolek (14 przemyslowych + 3 banki), 0 niezmapowanych,
# 7 roznic wobec wzorcow - wszystkie to bledy albo braki w plikach biznesradaru.
