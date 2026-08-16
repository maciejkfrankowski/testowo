# -*- coding: utf-8 -*-
"""Eksport slownika do CSV, zeby git pokazywal sensowne roznice.

    python slownik_csv.py

slownik.xlsx to plik binarny - 'git diff' nic z niego nie powie. Ten skrypt zrzuca
kazdy arkusz do slownik_csv/*.csv. Odpal przed commitem, wtedy w historii widac
DOKLADNIE ktory alias sie zmienil.
"""
import csv
import os
import sys

import openpyxl

TU = os.path.dirname(os.path.abspath(__file__)) if "__file__" in globals() else os.getcwd()
SLOWNIK = os.path.join(TU, sys.argv[1] if len(sys.argv) > 1 else "slownik.xlsx")
KATALOG = os.path.join(TU, "slownik_csv")

if not os.path.exists(SLOWNIK):
    raise SystemExit(f"Brak pliku {SLOWNIK}")
os.makedirs(KATALOG, exist_ok=True)

wb = openpyxl.load_workbook(SLOWNIK, data_only=True)
for nazwa in wb.sheetnames:
    ws = wb[nazwa]
    sciezka = os.path.join(KATALOG, f"{nazwa}.csv")
    with open(sciezka, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            if all(c is None or str(c).strip() == "" for c in row):
                continue
            w.writerow(["" if c is None else c for c in row])
    print(f"  {nazwa:22} -> slownik_csv/{nazwa}.csv")
print("\nGotowe. Dodaj slownik_csv/ do commita razem ze slownik.xlsx.")
